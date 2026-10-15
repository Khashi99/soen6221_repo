#!/usr/bin/env python3
"""Build the measured Person C report with the bundled document runtime."""
import csv
import json
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, text, width in zip(t.rows[0].cells, headers, widths):
        cell.text = text
        cell.width = Inches(width)
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EEF4'); cell._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for data in rows:
        row = t.add_row()
        for cell, text, width in zip(row.cells, data, widths):
            cell.text = str(text)
            cell.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)
    for i, row in enumerate(t.rows):
        cant_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant_split)
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge); e.set(qn('w:w'), '85'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            cell._tc.get_or_add_tcPr().append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 or len(headers) == 4 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    r.bold = i == 0 or data_is_total(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def data_is_total(row):
    return row.cells[0].text == 'TOTAL'


def main():
    manifest = json.loads((ROOT / 'results/run_manifest.json').read_text())
    config = json.loads((ROOT / 'patterns.independent.json').read_text())
    snapshots = json.loads((ROOT / 'snapshots.json').read_text())
    with (ROOT / 'results/sensitivity.csv').open() as f:
        measured = list(csv.DictReader(f))
    by = {(r['framework'], r['variant']): r for r in measured}
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.7)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        st = doc.styles[name]
        st.font.name = 'Calibri'; st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_after = Pt(7)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.05
    doc.styles['Title'].font.size = Pt(20)
    doc.styles['Heading 1'].font.size = Pt(14)
    doc.core_properties.title = 'SOEN 6221 Person C SATD keyword validation'
    doc.core_properties.subject = 'Fresh independent comparison and human review preparation'
    doc.core_properties.author = ''
    doc.add_paragraph('SOEN 6221 Person C SATD keyword validation', 'Title')
    doc.add_paragraph('Fresh independent study  |  13 October 2026', 'Subtitle')
    doc.styles['Subtitle'].font.color.rgb = RGBColor(0, 0, 0)
    # Strip inherited title/paragraph borders from the bundled default template.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    for border in doc.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    doc.add_paragraph('The proposed detector found 641 strict candidates and 693 full-set candidates in 66,754 Python comment blocks across ten pinned framework repositories. The 100-comment review sample is ready. The pattern set is provisional until human validation and original-paper source checks are complete.')
    doc.add_paragraph('Proposed keyword configuration', 'Heading 1')
    keyword_rows = [
        ['S01 to S04', 'TODO FIXME XXX HACK', 'Strict', 'Correction or shortcut markers [1, 2]'],
        ['P01', 'workaround variants', 'Loose', 'Compensating compromise [1, 3]'],
        ['P02', 'hardcode variants', 'Loose', 'Fixed-value shortcut [1, 3]'],
        ['P03', 'temporary solution', 'Loose', 'Provisional implementation [1, 3]'],
        ['P04', 'kludge', 'Loose', 'Admitted workaround [1, 3]'],
        ['P05', 'ugly', 'Loose', 'Admitted quality concern [1, 3]'],
        ['D01', 'DEPRECATED', 'Diagnostic', 'Routine deprecation may inflate counts [1]'],
    ]
    table(doc, ['IDs', 'Pattern terms', 'Use', 'Inclusion reason and lead'], keyword_rows, [.8, 1.8, .9, 3.6])
    doc.add_paragraph('Strict markers use case-insensitive whole-word regexes. The supplied context describes MAT fuzzy matching; these exact rules are an adaptation. The five loose phrase families are study proposals motivated by that context, not a verified transcription of Vidoni or the Potdar–Shihab 62-pattern dictionary. Exact regexes, flags and provenance are saved in patterns.independent.json and keyword_configuration.csv.')
    doc.add_paragraph('Extraction and comparison rules', 'Heading 1')
    doc.add_paragraph('Analyze tracked .py files, including tests and Python examples, from each new pinned snapshot. Python tokenize extracts actual # comments and excludes docstrings. Merge adjacent standalone comments only at equal indentation, without crossing blank lines or code. Inline comments remain separate. Keep raw text and source context; remove delimiters and leading whitespace only from the matching copy. Exclude the listed vendor, environment and cache directories; source packages named build or dist are retained. No tokenization failures occurred.')
    doc.add_paragraph('Count each block once per variant. Anchoring applies to the start of every constituent comment line and is independent of word boundaries. Standalone DEPRECATED is omitted from the full set; comments containing both DEPRECATED and a retained debt pattern can still match. Density uses all extracted blocks, retaining identical text at distinct locations.')

    doc.add_page_break()
    doc.add_paragraph('Measured sensitivity comparison', 'Heading 1')
    rows = []
    variants = ['strict_plus_deprecated', 'strict_without_deprecated', 'anchored_strict_without_deprecated', 'full_without_deprecated']
    for f in config['frameworks'] + ['TOTAL']:
        values = [by[(f, v)] for v in variants]
        rows.append([f, f"{int(values[0]['comment_blocks']):,}"] + [r['candidates'] for r in values])
    table(doc, ['Framework', 'Blocks', 'Strict plus D', 'Strict no D', 'Anchored', 'Full no D'], rows, [1.9, 1.0, 1.05, 1.05, 1.05, 1.05])
    doc.add_paragraph('D denotes the standalone DEPRECATED diagnostic. Relative to strict without D, adding D introduces 972 candidates; anchoring removes 44; the proposed loose phrases add 52. Total densities per 1,000 blocks are 24.16, 9.60, 8.94 and 10.38, respectively. The CSV retains per-framework density and added/removed counts. Strict candidate coverage is 641/693 = 92.50%; this is not SATD recall.')
    doc.add_paragraph('Manual validation preparation', 'Heading 1')
    doc.add_paragraph('With seed 6221, sample without replacement from two disjoint pools: 50 of 641 strict candidates and 50 of 52 loose-only candidates. Inclusion probabilities are 7.8003% and 96.1538%. All ten frameworks are represented, with 1 to 30 comments per framework; this is insufficient for reliable per-framework precision. Detector group is hidden from reviewer sheets and retained in the separate sampling administration file.')
    doc.add_paragraph('Reviewer 1 labels all 100 comments. A second human independently labels the designated random 30-comment overlap before discussion or adjudication. Use SATD, non-SATD or uncertain and give a reason. SATD requires a developer acknowledgement of an existing compromise, deficiency, shortcut or deferred corrective work. Context should distinguish optional features, resolved issues, quoted markers, negations, tool identifiers and expected test failures. A separate 23-row targeted edge-case audit is not part of the random precision sample.')
    doc.add_paragraph('No human labels have been supplied. Precision, Wilson intervals and Cohen’s kappa are pending. Once labels are complete, weight group precision by population sizes 641/693 and 52/693; do not average the equally sized samples. Preserve independent labels, use adjudication.csv for final decisions, report uncertainty bounds and paired binary N, and exclude uncertain pairs from kappa. Team targets of precision above 60% and kappa above 0.7 are evaluation goals, not universal thresholds.')

    doc.add_page_break()
    doc.add_paragraph('Reproducibility and scope', 'Heading 1')
    table(doc, ['Framework', 'Commit prefix'], [[r['framework'], r['commit'][:12]] for r in snapshots['repositories']], [3.6, 3.5])
    doc.add_paragraph('snapshots.json records full SHAs and acquisition times. The run and extraction manifests record detector 0.2.0, Python ' + manifest['python_version'] + ', seed 6221, hashes and scope. Fifteen synthetic tests passed. The package includes the corpus, matches and executable scripts; framework dependencies are unnecessary for tokenization.')
    doc.add_paragraph('The pilot’s reported 62,905 blocks, 640 strict and 1,846 full candidates remain unverified. New snapshots and patterns prevent direct comparison; similar strict totals do not establish reproduction. Reconcile Person B’s documentation before adopting the new pool. The later 600-candidate taxonomy study and Fleiss’ kappa are separate from this binary validation and Cohen’s kappa.')
    doc.add_paragraph('References and source status', 'Heading 1')
    doc.add_paragraph('[1] Supplied Person C project context and satd-keyword-validation skill, retained in references/. These supply task requirements and literature leads, not verified dictionaries.')
    doc.add_paragraph('[2] Guo et al. MAT A simple yet strong baseline for identifying self-admitted technical debt. arXiv:1910.13238. https://arxiv.org/abs/1910.13238. Source verification pending; no benchmark figures are used.')
    doc.add_paragraph('[3] Vidoni. Self-Admitted Technical Debt in R Packages An Exploratory Study. MSR 2021. DOI 10.1109/MSR52588.2021.00030. The context points to Section IV-C, p.185: https://melvidoni.rbind.io/publication/2021-rsatd-msr.pdf. Source verification pending; R findings do not validate this Python detector.')
    doc.add_paragraph('Limitations', 'Heading 1')
    doc.add_paragraph('Candidate counts measure detector sensitivity, not confirmed debt. Human labels, agreement and original-paper verification are outstanding. The proposed loose set is narrower than the unavailable pilot dictionary; pattern selection remains provisional and must not be tuned repeatedly to pass team targets. Any revision after calibration needs a fresh held-out validation sample. Reviewing flagged comments alone does not estimate absolute recall. Source hosts returned proxy access denials; their proposed network additions are saved but runtime application is unconfirmed.')
    output = ROOT / 'results/person_c_keyword_validation.docx'
    doc.save(output)
    print(output)


if __name__ == '__main__':
    main()
